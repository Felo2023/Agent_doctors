"""Synthetic-only source adapters for Word, Excel, and CSV.

These preserve source locations and are not medical data abstraction.
Unsupported/missing content is flagged instead of silently flattened.
"""

import csv
import hashlib
import json
from datetime import date, datetime, time
from pathlib import Path
import subprocess
import sys
import tempfile
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from openpyxl import load_workbook

from intake.extract import MAX_FILE_SIZE, check_fixture, checksum


MAX_CELLS = 10000
MAX_ZIP_BYTES = 100 * 1024 * 1024


def source_document(path, fixture_root, suffixes):
    path = check_fixture(path, fixture_root)
    if not path.is_file() or path.suffix.lower() not in suffixes:
        raise ValueError("unsupported source format")
    if path.stat().st_size > MAX_FILE_SIZE:
        raise ValueError("file exceeds prototype size limit")
    digest = checksum(path)
    return path, {"source_id": "SRC-" + digest[:16], "filename": path.name, "sha256": digest,
                  "format": path.suffix.lower().lstrip("."), "blocks": [], "warnings": [], "candidates": []}


def check_zip(path):
    with ZipFile(path) as archive:
        names = [item.filename for item in archive.infolist()]
        if "[Content_Types].xml" not in names or "_rels/.rels" not in names:
            raise ValueError("file is not a valid Office Open XML package")
        if (len(names) > 10000 or sum(item.file_size for item in archive.infolist()) > MAX_ZIP_BYTES
                or any(item.file_size > MAX_FILE_SIZE for item in archive.infolist())):
            raise ValueError("office archive exceeds extraction limits")
        if any(item.external_attr >> 16 & 0o170000 == 0o120000 for item in archive.infolist()):
            raise ValueError("office archive contains symlinks")
        return names


def vision_image(blob, suffix):
    """OCR embedded raster image locally; never send the bytes to a provider."""
    if sys.platform != "darwin":
        raise RuntimeError("local Apple Vision OCR requires macOS")
    with tempfile.TemporaryDirectory(prefix="synthetic-doc-image-") as temp:
        file = Path(temp) / ("embedded." + suffix)
        file.write_bytes(blob)
        result = subprocess.run(["swift", str(Path(__file__).with_name("vision_ocr.swift")), str(file)],
                                capture_output=True, text=True, timeout=120)
    if result.returncode:
        raise RuntimeError("embedded-image OCR unavailable: " + result.stderr.strip()[-300:])
    return json.loads(result.stdout)["lines"]


def extract_docx(path, *, fixture_root=None, ocr=vision_image):
    path, output = source_document(path, fixture_root, {".docx"})
    names = check_zip(path)
    if "word/comments.xml" in names:
        output["warnings"].append("comments present but not extracted")
    if "word/footnotes.xml" in names or "word/endnotes.xml" in names:
        output["warnings"].append("footnotes/endnotes present but not extracted")
    if any(name.startswith("word/charts/") for name in names):
        output["warnings"].append("Word charts present but not extracted")
    with ZipFile(path) as archive:
        body_xml = archive.read("word/document.xml")
        if b"<w:ins" in body_xml or b"<w:del" in body_xml:
            output["warnings"].append("tracked changes present; extraction may omit revisions")
        if b"txbxContent" in body_xml:
            output["warnings"].append("text boxes present; contents may be omitted")
    doc = Document(path)
    source_id = output["source_id"]

    def add(text, kind, locator):
        block_id = f"{source_id}:" + ":".join(str(value) for value in locator.values())
        output["blocks"].append({"id": block_id, "type": kind, "text": text,
                                  "locator": locator})

    def embedded(paragraph, prefix):
        for run_index, run in enumerate(paragraph.runs, 1):
            for image_index, blip in enumerate(run._r.xpath(".//a:blip"), 1):
                relationship = blip.get(qn("r:embed"))
                if not relationship or relationship not in doc.part.related_parts:
                    output["warnings"].append("unresolved embedded image reference at " + prefix)
                    continue
                part = doc.part.related_parts[relationship]
                mime = part.content_type
                suffix = {"image/png": "png", "image/jpeg": "jpg", "image/tiff": "tiff"}.get(mime)
                locator = {"paragraph": prefix, "run": run_index, "image": image_index}
                blob = part.blob
                ref = hashlib.sha256(blob).hexdigest()
                add("Embedded image (" + mime + ")", "image", {**locator, "sha256": ref[:16]})
                if suffix is None:
                    output["warnings"].append("embedded image type needs review: " + mime)
                    continue
                try:
                    lines = ocr(blob, suffix)
                    for line_index, line in enumerate(lines, 1):
                        output["blocks"].append({
                            "id": f"{source_id}:{prefix}:r{run_index}:image{image_index}:ocr:{line_index}",
                            "type": "ocr_line", "text": line["text"],
                            "locator": {**locator, "ocr_line": line_index, "bbox": line.get("bbox"),
                                        "ocr_confidence": line.get("confidence")},
                        })
                except (RuntimeError, ValueError, subprocess.TimeoutExpired, OSError) as exc:
                    output["warnings"].append("embedded image OCR failed at " + prefix + ": " + type(exc).__name__)

    for index, element in enumerate(doc.element.body.iterchildren(), 1):
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, doc)
            locator = {"body_element": index, "paragraph": 1}
            if paragraph.text.strip():
                add(paragraph.text, "heading" if paragraph.style.name.startswith("Heading") else "paragraph", locator)
            embedded(paragraph, f"body{index}")
        elif element.tag == qn("w:tbl"):
            table = Table(element, doc)
            for row_number, row in enumerate(table.rows, 1):
                for column_number, cell in enumerate(row.cells, 1):
                    locator = {"body_element": index, "row": row_number, "column": column_number}
                    add(cell.text, "table_cell", locator)
                    for paragraph_number, paragraph in enumerate(cell.paragraphs, 1):
                        embedded(paragraph, f"body{index}:r{row_number}:c{column_number}:p{paragraph_number}")
        elif element.tag != qn("w:sectPr"):
            output["warnings"].append("unsupported Word body element at index " + str(index))
    for number, section in enumerate(doc.sections, 1):
        for kind, element in (("header", section.header), ("footer", section.footer)):
            for para_index, paragraph in enumerate(element.paragraphs, 1):
                if paragraph.text.strip():
                    add(paragraph.text, kind, {"section": number, kind: para_index})
                embedded(paragraph, f"section{number}:{kind}{para_index}")
    return output


def serialized(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def extract_xlsx(path, *, fixture_root=None):
    path, output = source_document(path, fixture_root, {".xlsx"})
    names = check_zip(path)
    if any(name.startswith("xl/externalLinks/") for name in names):
        output["warnings"].append("external workbook links are not resolved")
    if any(name.startswith(("xl/charts/", "xl/drawings/", "xl/pivotTables/")) for name in names):
        output["warnings"].append("charts, drawings or pivot tables present; their content is not extracted")
    if any("comment" in name.lower() for name in names):
        output["warnings"].append("spreadsheet comments present but not extracted")
    with ZipFile(path) as archive:
        if any(b"<mergeCells" in archive.read(name) for name in names if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")):
            output["warnings"].append("merged cells present; table layout needs review")
    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    count = 0
    try:
        for sheet in formulas.worksheets:
            cached = values[sheet.title]
            if sheet.max_row > MAX_CELLS or sheet.max_column > 256:
                raise ValueError("worksheet dimensions exceed synthetic extraction limits")
            if sheet.sheet_state != "visible":
                output["warnings"].append("hidden worksheet included: " + sheet.title)
            output.setdefault("sheets", []).append({"name": sheet.title, "state": sheet.sheet_state})
            for row_formula, row_cached in zip(sheet.iter_rows(), cached.iter_rows()):
                for cell, saved in zip(row_formula, row_cached):
                    if cell.value is None and saved.value is None:
                        continue
                    count += 1
                    if count > MAX_CELLS:
                        raise ValueError("workbook exceeds synthetic populated-cell limit")
                    formula = cell.value if cell.data_type == "f" else None
                    if formula is not None and saved.value is None:
                        output["warnings"].append("formula has no saved value: " + sheet.title + "!" + cell.coordinate)
                    output["blocks"].append({
                        "id": f"{output['source_id']}:sheet:{sheet.title}:cell:{cell.coordinate}",
                        "type": "cell",
                        "text": str(serialized(cell.value)),
                        "locator": {"sheet": sheet.title, "cell": cell.coordinate},
                        "value": serialized(cell.value) if formula is None else None,
                        "formula": formula, "cached_value": serialized(saved.value) if formula is not None else None,
                        "number_format": cell.number_format,
                    })
    finally:
        formulas.close()
        values.close()
    return output


def column_name(number):
    text = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        text = chr(65 + remainder) + text
    return text


def extract_csv(path, *, fixture_root=None):
    path, output = source_document(path, fixture_root, {".csv", ".tsv"})
    # Reject undecodable input; a guessed encoding could silently alter a lab value.
    text = path.read_text(encoding="utf-8-sig")
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    if path.suffix.lower() == ".csv":
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            delimiter = dialect.delimiter
        except csv.Error:
            output["warnings"].append("CSV delimiter could not be inferred; comma used")
    output["delimiter"] = delimiter
    reader = csv.reader(text.splitlines(keepends=True), delimiter=delimiter)
    count = 0
    for row_number, row in enumerate(reader, 1):
        for column_number, value in enumerate(row, 1):
            if not value:
                continue
            count += 1
            if count > MAX_CELLS:
                raise ValueError("CSV exceeds synthetic populated-cell limit")
            coordinate = f"{column_name(column_number)}{row_number}"
            output["blocks"].append({
                "id": f"{output['source_id']}:row:{row_number}:cell:{coordinate}",
                "type": "cell", "text": value,
                "locator": {"row": row_number, "column": column_number, "cell": coordinate},
                "value": value, "formula": None, "cached_value": None,
            })
    return output
