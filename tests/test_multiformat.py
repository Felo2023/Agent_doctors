import importlib.util
import json
from pathlib import Path
import re
import sys
import tempfile
import unittest
from zipfile import ZipFile


HAS_DEPS = all(importlib.util.find_spec(name) for name in ("pymupdf", "docx", "openpyxl"))
if HAS_DEPS:
    import pymupdf
    from docx import Document
    from docx.shared import Inches
    from openpyxl import Workbook
    from intake.extract import extract_snapshot
    from intake.markdown import write_bundle
    from intake.corrections import propose, review


@unittest.skipUnless(HAS_DEPS, "install requirements-intake.txt")
class MultiFormatTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="synthetic-multiformat-")
        self.root = Path(self.temp.name)
        (self.root / ".synthetic-fixture").write_text("SYNTHETIC_INTAKE_FIXTURE_V1")

    def tearDown(self):
        self.temp.cleanup()

    def word(self, with_image=False):
        document = Document()
        document.add_heading("Fictional laboratory summary", 1)
        document.add_paragraph("Only synthetic observations appear here.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Prueba"
        table.cell(0, 1).text = "Resultado"
        table.cell(1, 0).text = "Creatinina"
        table.cell(1, 1).text = "1,4 mg/dL"
        if with_image:
            image_doc = pymupdf.open()
            page = image_doc.new_page(width=600, height=200)
            page.insert_text((25, 55), "Glucosa 92 mg/dL", fontsize=22)
            from io import BytesIO
            document.add_picture(BytesIO(page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False).tobytes("png")), width=Inches(4))
            image_doc.close()
        path = self.root / "fictional-report.docx"
        document.save(path)
        return path

    def workbook(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Results"
        sheet["A1"] = "Value"
        sheet["A2"] = 10
        sheet["B2"] = "=A2*2"
        hidden = book.create_sheet("Old results")
        hidden.sheet_state = "hidden"
        hidden["A1"] = "unreviewed historical value"
        path = self.root / "fictional-results.xlsx"
        book.save(path)
        return path

    def test_word_body_table_and_markdown_keep_source(self):
        path = self.word()
        snapshot = extract_snapshot([path], self.root / "snapshot.json", fixture_root=self.root)
        doc = snapshot["documents"][0]
        self.assertEqual(doc["format"], "docx")
        self.assertEqual(len([b for b in doc["blocks"] if b["type"] == "table_cell"]), 4)
        self.assertTrue(any("Fictional laboratory" in b["text"] for b in doc["blocks"]))
        files = write_bundle(snapshot, self.root, fixture_root=self.root)
        markdown = (self.root / (doc["source_id"] + ".md")).read_text()
        self.assertIn("Creatinina", markdown)
        self.assertIn(doc["blocks"][0]["id"], markdown)
        self.assertEqual(len(files), 2)
        self.assertTrue((self.root / "index.md").exists())

    @unittest.skipUnless(sys.platform == "darwin", "local Vision OCR requires macOS")
    def test_word_embedded_image_is_ocr_source_block(self):
        doc = extract_snapshot([self.word(with_image=True)], self.root / "snapshot.json", fixture_root=self.root)["documents"][0]
        self.assertTrue(any(b["type"] == "image" for b in doc["blocks"]))
        self.assertTrue(any(b["type"] == "ocr_line" and "Glucosa" in b["text"] for b in doc["blocks"]))

    def test_workbook_formula_cache_absence_hidden_sheet_and_markdown(self):
        doc = extract_snapshot([self.workbook()], self.root / "snapshot.json", fixture_root=self.root)["documents"][0]
        formula = next(b for b in doc["blocks"] if b["locator"].get("cell") == "B2")
        self.assertEqual(formula["formula"], "=A2*2")
        self.assertIsNone(formula["cached_value"])
        self.assertTrue(any("no saved value" in w for w in doc["warnings"]))
        self.assertEqual(next(s for s in doc["sheets"] if s["name"] == "Old results")["state"], "hidden")
        snap = {"case_id": "SYNTH-001", "documents": [doc], "conflicts": []}
        write_bundle(snap, self.root, fixture_root=self.root)
        markdown = (self.root / (doc["source_id"] + ".md")).read_text()
        self.assertIn("=A2*2", markdown)
        self.assertIn("not available", markdown)
        self.assertIn("Old results (hidden)", markdown)

    def test_formula_and_saved_result_are_distinct(self):
        path = self.workbook()
        rewritten = self.root / "fictional-with-cache.xlsx"
        with ZipFile(path) as original, ZipFile(rewritten, "w") as updated:
            for member in original.infolist():
                data = original.read(member.filename)
                if member.filename == "xl/worksheets/sheet1.xml":
                    data, count = re.subn(rb"(<f>A2\*2</f><v>)[^<]*(</v>)", rb"\g<1>20\2", data)
                    self.assertEqual(count, 1)
                updated.writestr(member, data)
        doc = extract_snapshot([rewritten], self.root / "snapshot.json", fixture_root=self.root)["documents"][0]
        result = next(b for b in doc["blocks"] if b["locator"].get("cell") == "B2")
        self.assertEqual(result["formula"], "=A2*2")
        self.assertEqual(result["cached_value"], 20)

    def test_csv_escaping_preserves_exact_json_and_unsupported_format_fails(self):
        path = self.root / "fictional-values.csv"
        path.write_text('Name,Value\n"<review>","2|3"\n', encoding="utf-8")
        snap = extract_snapshot([path], self.root / "snapshot.json", fixture_root=self.root)
        doc = snap["documents"][0]
        self.assertTrue(any(b["value"] == "<review>" for b in doc["blocks"]))
        write_bundle(snap, self.root, fixture_root=self.root)
        md = (self.root / (doc["source_id"] + ".md")).read_text()
        self.assertIn("&lt;review&gt;", md)
        self.assertIn("2\\|3", md)
        unsupported = self.root / "later.pptx"
        unsupported.write_bytes(b"synthetic")
        with self.assertRaisesRegex(ValueError, "unsupported"):
            extract_snapshot([unsupported], self.root / "ignored.json", fixture_root=self.root)
        snap["documents"][0]["source_id"] = "../../outside"
        with self.assertRaisesRegex(ValueError, "invalid source ID"):
            write_bundle(snap, self.root, fixture_root=self.root)

    def test_corrections_are_separate_and_reviewed_without_mutating_source(self):
        snapshot = extract_snapshot([self.word()], self.root / "snapshot.json", fixture_root=self.root)
        doc = snapshot["documents"][0]
        write_bundle(snapshot, self.root, fixture_root=self.root)
        original = (self.root / (doc["source_id"] + ".md")).read_bytes()
        target = next(b for b in doc["blocks"] if b["text"] == "1,4 mg/dL")
        ledger = self.root / "corrections.jsonl"
        event = propose(snapshot, ledger, fixture_root=self.root, source_id=doc["source_id"],
                        block_id=target["id"], suggested_text="1,8 mg/dL", reason="example reviewer adjustment", author="synthetic-reviewer")
        self.assertEqual(event["status"], "pending")
        decision = review(snapshot, ledger, fixture_root=self.root, proposal_id=event["event_id"],
                          decision="accepted", reviewer="synthetic-second-reviewer", reason="verified against fictional source")
        self.assertEqual(decision["decision"], "accepted")
        self.assertEqual(len(ledger.read_text().splitlines()), 2)
        correction_view = (self.root / "corrections.md").read_text()
        self.assertIn("Status: **accepted**", correction_view)
        self.assertIn(doc["source_id"] + ".md#block-", correction_view)
        self.assertEqual((self.root / (doc["source_id"] + ".md")).read_bytes(), original)
        write_bundle(snapshot, self.root, fixture_root=self.root)
        self.assertEqual((self.root / (doc["source_id"] + ".md")).read_bytes(), original)
        with self.assertRaisesRegex(ValueError, "already reviewed"):
            review(snapshot, ledger, fixture_root=self.root, proposal_id=event["event_id"],
                   decision="rejected", reviewer="synthetic-second-reviewer", reason="duplicate")
        with self.assertRaisesRegex(ValueError, "existing source block"):
            propose(snapshot, ledger, fixture_root=self.root, source_id=doc["source_id"],
                    block_id="invented", suggested_text="x", reason="y", author="z")


if __name__ == "__main__":
    unittest.main()
