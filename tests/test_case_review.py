import importlib.util
from pathlib import Path
import tempfile
import unittest
import threading
from http.server import ThreadingHTTPServer
import json
import urllib.error
import urllib.request


if importlib.util.find_spec("pymupdf"):
    import pymupdf
    from docx import Document
    from openpyxl import Workbook
    from case_review.store import ConflictError, ReviewStore
    from bus.prototype import Bus, Handler as MailHandler
    from intake.corrections import propose, review
    from intake.extract import extract_snapshot
    from intake.markdown import write_bundle
else:
    pymupdf = None


@unittest.skipUnless(pymupdf, "install requirements-intake.txt to run case review tests")
class ReviewedSyntheticCaseTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="synthetic-review-")
        self.root = Path(self.temp.name)
        (self.root / ".synthetic-fixture").write_text("SYNTHETIC_INTAKE_FIXTURE_V1")

    def tearDown(self):
        self.temp.cleanup()

    def pdf(self, name, lab):
        document = pymupdf.open()
        page = document.new_page(width=600, height=800)
        page.insert_text((40, 80), "Fecha: 2025-12-04", fontsize=15)
        page.insert_text((40, 110), lab, fontsize=15)
        file = self.root / name
        document.save(file)
        document.close()
        return file

    def snapshot(self, paths):
        snap = extract_snapshot(paths, self.root / "extraction.json", fixture_root=self.root)
        write_bundle(snap, self.root, fixture_root=self.root)
        return snap

    def test_conflicts_block_commit_until_each_candidate_reviewed(self):
        self.snapshot([self.pdf("one.pdf", "Glucosa 92 mg/dL"), self.pdf("two.pdf", "Glucosa 128 mg/dL")])
        store = ReviewStore(self.root)
        try:
            ids = list(store.candidates)
            with self.assertRaisesRegex(ValueError, "review all"):
                store.commit(0)
            store.decide(ids[0], "verify", "reviewer-1", "checked synthetic original", 0)
            with self.assertRaisesRegex(ValueError, "review all"):
                store.commit(1)
            store.decide(ids[1], "verify", "reviewer-1", "checked synthetic original", 1)
            with self.assertRaisesRegex(ValueError, "conflicting verified"):
                store.commit(2)
            store.decide(ids[1], "exclude", "reviewer-2", "synthetic duplicate with discrepant value", 2)
            first = store.commit(3)
            self.assertEqual(first["version"], 1)
            self.assertEqual(first["observations"][0]["value_decimal"], "92")
            self.assertFalse(first["ready_for_diagnosis"])
            self.assertEqual(first["held_candidates"][0]["status"], "exclude")
            with self.assertRaisesRegex(ConflictError, "reload"):
                store.decide(ids[0], "defer", "reviewer-1", "stale screen", 3)
            with self.assertRaisesRegex(ValueError, "no new reviewed decisions"):
                store.commit(4)
            store.decide(ids[1], "defer", "reviewer-2", "awaiting synthetic source clarification", 4)
            second = store.commit(5)
            self.assertEqual(second["version"], 2)
            self.assertEqual(second["previous_hash"], first["hash"])
            self.assertEqual(store.version(1)["held_candidates"][0]["status"], "exclude")
            self.assertEqual(store.version(2)["held_candidates"][0]["status"], "defer")
        finally:
            store.close()
        reopened = ReviewStore(self.root)
        try:
            self.assertEqual(reopened.version()["version"], 2)
        finally:
            reopened.close()

    def test_accepted_correction_still_requires_candidate_verification(self):
        snap = self.snapshot([self.pdf("one.pdf", "Creatinina 1,4 mg/dL")])
        doc = snap["documents"][0]
        item = doc["candidates"][0]
        ledger = self.root / "corrections.jsonl"
        proposal = propose(snap, ledger, fixture_root=self.root, source_id=doc["source_id"],
                           block_id=item["source"]["block_id"], suggested_text="Creatinina 1,8 mg/dL",
                           reason="fictional image clarification", author="reviewer-a")
        store = ReviewStore(self.root)
        try:
            with self.assertRaisesRegex(ValueError, "accepted correction"):
                store.decide(item["id"], "verify", "reviewer-b", "checked fictional report", 0,
                             correction_id=proposal["event_id"])
            self.assertEqual(store.state()["revision"], 0)
            review(snap, ledger, fixture_root=self.root, proposal_id=proposal["event_id"],
                   decision="accepted", reviewer="reviewer-b", reason="verified fictional page")
            store.decide(item["id"], "verify", "reviewer-b", "checked fictional page", 0,
                         correction_id=proposal["event_id"])
            self.assertEqual(store.state()["case_version"], 0)  # accepted correction did not auto-commit
            version = store.commit(1)
            observation = version["observations"][0]
            self.assertEqual(observation["value_decimal"], "1.8")
            self.assertEqual(observation["original_value"], "1,4")
            self.assertEqual(observation["source"]["exact_text"], "Creatinina 1,4 mg/dL")
            self.assertEqual(observation["correction_id"], proposal["event_id"])
        finally:
            store.close()

    def test_source_change_and_snapshot_change_fail_closed(self):
        file = self.pdf("one.pdf", "Glucosa 92 mg/dL")
        self.snapshot([file])
        store = ReviewStore(self.root)
        try:
            candidate_id = next(iter(store.candidates))
            file.write_bytes(file.read_bytes() + b"tampering")
            with self.assertRaisesRegex(ConflictError, "original source changed"):
                store.decide(candidate_id, "verify", "reviewer", "checked original", 0)
            self.assertEqual(store.state()["revision"], 0)
        finally:
            store.close()

    def test_extraction_snapshot_change_blocks_review(self):
        file = self.pdf("one.pdf", "Glucosa 92 mg/dL")
        self.snapshot([file])
        store = ReviewStore(self.root)
        try:
            candidate_id = next(iter(store.candidates))
            with (self.root / "extraction.json").open("a") as extraction:
                extraction.write(" ")
            with self.assertRaisesRegex(ConflictError, "extraction changed"):
                store.decide(candidate_id, "verify", "reviewer", "checked fictional PDF", 0)
            with self.assertRaisesRegex(ConflictError, "extraction changed"):
                store.commit(0)
            self.assertEqual(store.state()["revision"], 0)
        finally:
            store.close()

    def test_word_history_requires_independent_review_and_retains_timeline_provenance(self):
        word = self.root / "fictional-history.docx"
        document = Document()
        document.add_paragraph("Patient says fatigue started around early 2024.")
        document.save(word)
        snap = extract_snapshot([word], self.root / "extraction.json", fixture_root=self.root,
                                case_id="SYNTH-002")
        write_bundle(snap, self.root, fixture_root=self.root)
        store = ReviewStore(self.root)
        try:
            self.assertEqual(store.case_id, "SYNTH-002")
            block = snap["documents"][0]["blocks"][0]
            proposal = store.propose_fact(snap["documents"][0]["source_id"], block["id"],
                                          "symptom", "Fatigue began in early 2024 (reported by patient)",
                                          "2024", "approximate", "patient_reported", "proposer-A", 0)
            self.assertEqual(store.state()["pending_facts"], 1)
            with self.assertRaisesRegex(ValueError, "independent reviewer"):
                store.decide_fact(proposal["fact_id"], "verify", "proposer-A", "same person", 1)
            with self.assertRaisesRegex(ValueError, "review all proposed"):
                store.commit(1)
            store.decide_fact(proposal["fact_id"], "verify", "reviewer-B",
                              "Read original fictional Word paragraph", 1)
            version = store.commit(2)
            self.assertEqual(version["observations"], [])
            self.assertEqual(version["source_statements"][0]["original_text"], block["text"])
            self.assertEqual(version["source_statements"][0]["attribution"], "patient_reported")
            self.assertEqual(version["timeline"][0]["precision"], "approximate")
            self.assertEqual(version["timeline"][0]["date"], "2024")
            self.assertFalse(version["ready_for_diagnosis"])
        finally:
            store.close()

    def test_formula_without_saved_value_cannot_be_promoted_as_statement(self):
        sheet = Workbook()
        sheet.active["A1"] = "=1+1"
        file = self.root / "fictional.xlsx"
        sheet.save(file)
        snap = extract_snapshot([file], self.root / "extraction.json", fixture_root=self.root)
        store = ReviewStore(self.root)
        try:
            block = snap["documents"][0]["blocks"][0]
            with self.assertRaisesRegex(ValueError, "formula without a saved value"):
                store.propose_fact(snap["documents"][0]["source_id"], block["id"], "test_report",
                                   "Measured value was 2", None, "unknown", "source_text", "reviewer-A", 0)
        finally:
            store.close()

    def test_case_bound_mail_is_pinned_to_reviewed_version(self):
        snap = extract_snapshot([self.pdf("fictional.pdf", "Glucosa 92 mg/dL")],
                                self.root / "extraction.json", fixture_root=self.root,
                                case_id="SYNTH-002")
        store = ReviewStore(self.root)
        try:
            cid = snap["documents"][0]["candidates"][0]["id"]
            store.decide(cid, "verify", "reviewer-1", "checked fictional original", 0)
            version = store.commit(1)
            bus = Bus(str(self.root / "mail.sqlite"), "z" * 32,
                      server_passwords={"doctor-a": "a", "doctor-b": "b"},
                      case_id="SYNTH-002", case_root=self.root)
            try:
                bus.register("doctor-a", "http://127.0.0.1:44101", "ses_A",
                             "anthropic/claude-opus-5-5", "a" * 40)
                bus.register("doctor-b", "http://127.0.0.1:44102", "ses_B",
                             "openai/gpt-6.1-sol", "b" * 40)
                bus.set_phase("discussion")
                self.assertEqual(bus.read_case()["hash"], version["hash"])
                with self.assertRaisesRegex(ValueError, "case version"):
                    bus.send("doctor-a", "ses_A", "doctor-b", "stale", "one",
                             case_id="SYNTH-002", case_version=2)
                self.assertEqual(bus.send("doctor-a", "ses_A", "doctor-b", "read source", "two",
                                          case_id="SYNTH-002", case_version=1)["state"], "queued")
                document = snap["documents"][0]
                date_block = document["blocks"][0]
                fact = store.propose_fact(document["source_id"], date_block["id"], "test_report",
                                          "Fictional report date was 2025-12-04", "2025-12-04", "exact",
                                          "source_text", "proposer-1", 2)
                store.decide_fact(fact["fact_id"], "verify", "reviewer-2", "checked original date line", 3)
                store.commit(4)
                with self.assertRaisesRegex(ValueError, "bound case version changed"):
                    bus.read_case()
                with self.assertRaisesRegex(ValueError, "bound case version changed"):
                    bus.send("doctor-a", "ses_A", "doctor-b", "untrusted", "three",
                             case_id="SYNTH-002", case_version=1)
            finally:
                bus.case_store.close()
                bus.db.close()
        finally:
            store.close()

    def test_http_role_sees_only_bound_case_and_rejects_wrong_version(self):
        snap = extract_snapshot([self.pdf("fictional.pdf", "Glucosa 92 mg/dL")],
                                self.root / "extraction.json", fixture_root=self.root,
                                case_id="SYNTH-002")
        reviewer = ReviewStore(self.root)
        try:
            reviewer.decide(snap["documents"][0]["candidates"][0]["id"], "verify",
                            "reviewer-B", "checked fictional source", 0)
            reviewer.commit(1)
        finally:
            reviewer.close()
        bus = Bus(str(self.root / "bus.sqlite"), "z" * 32,
                  server_passwords={"doctor-a": "a", "doctor-b": "b"},
                  case_id="SYNTH-002", case_root=self.root)
        server = ThreadingHTTPServer(("127.0.0.1", 0), MailHandler)
        server.bus = bus
        bus.running = True
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            bus.register("doctor-a", "http://127.0.0.1:44101", "ses_A",
                         "anthropic/claude-opus-5-5", "a" * 40)
            bus.register("doctor-b", "http://127.0.0.1:44102", "ses_B",
                         "openai/gpt-6.1-sol", "b" * 40)
            bus.set_phase("discussion")
            url = f"http://127.0.0.1:{server.server_port}"

            def request(path, token, body=None):
                req = urllib.request.Request(url + path,
                    headers={"Authorization": "Bearer " + token, "Content-Type": "application/json"},
                    data=json.dumps(body).encode() if body is not None else None,
                    method="POST" if body is not None else "GET")
                try:
                    with urllib.request.urlopen(req, timeout=3) as response:
                        return response.status, json.loads(response.read())
                except urllib.error.HTTPError as error:
                    return error.code, json.loads(error.read())

            status, case = request("/case", "a" * 40)
            self.assertEqual(status, 200)
            self.assertEqual(case["case_id"], "SYNTH-002")
            self.assertEqual(case["version"], 1)
            self.assertFalse(case["ready_for_diagnosis"])
            self.assertEqual(request("/case", "foreign-token")[0], 401)
            letter = {"case_id": "SYNTH-002", "case_version": 1, "sender_session": "ses_A",
                      "recipient": "doctor-b", "body": "fictional source reference", "dedupe_key": "case-http"}
            self.assertEqual(request("/mail", "a" * 40, {**letter, "case_id": "SYNTH-001"})[0], 400)
            self.assertEqual(request("/mail", "a" * 40, {**letter, "case_version": 2})[0], 400)
            self.assertEqual(request("/mail", "b" * 40, letter)[0], 400)
            self.assertEqual(request("/mail", "a" * 40, letter)[0], 200)
            self.assertEqual(len(bus.list_mail()), 1)
        finally:
            bus.running = False
            server.shutdown()
            server.server_close()
            thread.join(timeout=3)
            bus.case_store.close()
            bus.db.close()


if __name__ == "__main__":
    unittest.main()
