import importlib.util
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile


if importlib.util.find_spec("docx"):
    from docx import Document
    from scripts.create_patient_history_form import build
else:
    Document = None


@unittest.skipUnless(Document, "install requirements-intake.txt for Word document checks")
class PatientHistoryFormTests(unittest.TestCase):
    def test_patient_form_is_editable_complete_and_not_prefilled_with_case_details(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "form.docx"
            build(path)
            with ZipFile(path) as archive:
                self.assertIsNone(archive.testzip())
            document = Document(path)
            headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
            self.assertEqual(len(headings), 11)
            self.assertTrue(any("Síntomas actuales" in value for value in headings))
            self.assertTrue(any("Medicamentos y alergias" in value for value in headings))
            self.assertTrue(any("Salud reproductiva (si corresponde)" in value for value in headings))
            self.assertGreater(len(document.tables), 30)
            cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            self.assertTrue(any("[Escriba aquí / Type here]" in value for value in cells))
            self.assertTrue(any("☐ Sí" in value for value in cells))
            self.assertTrue(any("no sé" in p.text for p in document.paragraphs))
            self.assertFalse(any("Gabriela" in value for value in headings + cells))
            self.assertAlmostEqual(document.tables[0].columns[0].width.inches, 2.37, places=1)
            self.assertAlmostEqual(document.tables[0].columns[1].width.inches, 4.59, places=1)


if __name__ == "__main__":
    unittest.main()
