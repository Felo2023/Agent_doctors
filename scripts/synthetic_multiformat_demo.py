"""Build fictional PDFs, DOCX, XLSX and CSV with source-linked Markdown.

The demo records a proposed correction separately; it never rewrites sources.
"""

from pathlib import Path
import tempfile

from docx import Document
from openpyxl import Workbook

from intake.corrections import propose
from intake.extract import extract_snapshot
from intake.markdown import write_bundle
from scripts.synthetic_intake_demo import create_page


def main():
    root = Path(tempfile.mkdtemp(prefix="synthetic-multiformat-"))
    root.chmod(0o700)
    (root / ".synthetic-fixture").write_text("SYNTHETIC_INTAKE_FIXTURE_V1", encoding="utf-8")
    pdf = root / "fictional-scan.pdf"
    create_page(pdf, "1,4", scan=True)

    word = root / "fictional-history.docx"
    document = Document()
    document.add_heading("Fictional history", 1)
    document.add_paragraph("This text is entirely synthetic.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Source"
    table.cell(0, 1).text = "Note"
    table.cell(1, 0).text = "Family report"
    table.cell(1, 1).text = "No real patient data"
    document.save(word)

    excel = root / "fictional-table.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Synthetic results"
    sheet["A1"] = "Value"
    sheet["A2"] = 10
    sheet["B2"] = "=A2*2"
    workbook.save(excel)

    csv = root / "fictional-cells.csv"
    csv.write_text("Label,Value\nSynthetic,92\n", encoding="utf-8")

    snapshot = extract_snapshot([pdf, word, excel, csv], root / "extraction.json", fixture_root=root)
    files = write_bundle(snapshot, root, fixture_root=root)
    word_record = next(doc for doc in snapshot["documents"] if doc["format"] == "docx")
    proposal = propose(snapshot, root / "corrections.jsonl", fixture_root=root,
                       source_id=word_record["source_id"], block_id=word_record["blocks"][0]["id"],
                       suggested_text="Fictional history (sample correction)",
                       reason="Synthetic demonstration of a separate correction record", author="demo-reviewer")
    print("Synthetic directory:", root)
    print("Formats:", [doc["format"] for doc in snapshot["documents"]])
    print("Markdown files:", len(files))
    print("Correction status:", proposal["status"], "(original Markdown unchanged)")


if __name__ == "__main__":
    main()
