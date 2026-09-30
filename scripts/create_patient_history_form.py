"""Create a patient-completable bilingual clinical history Word questionnaire.

The generator has no patient records and never reads the source-PDF directory.
Its default output is the reusable, unfilled template at the repository root.
Keep completed patient copies in a private/gitignored location.
"""

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = RGBColor(23, 51, 79)
TEAL = RGBColor(14, 108, 120)
GRAY = RGBColor(87, 100, 113)
INK = RGBColor(30, 39, 48)
FILL = "F2F6FA"
LINE = "CFDCE6"


def shade(cell, color=FILL):
    properties = cell._tc.get_or_add_tcPr()
    existing = properties.find(qn("w:shd"))
    if existing is None:
        existing = OxmlElement("w:shd")
        properties.append(existing)
    existing.set(qn("w:fill"), color)


def set_cell_border(cell):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for name in ("top", "bottom", "left", "right"):
        side = OxmlElement("w:" + name)
        side.set(qn("w:val"), "single")
        side.set(qn("w:sz"), "5")
        side.set(qn("w:color"), LINE)
        borders.append(side)


def no_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def repeating_header(row):
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(header)


def set_table_width(table, widths):
    table.autofit = False
    total = sum(widths)
    properties = table._tbl.tblPr
    element = properties.find(qn("w:tblW"))
    if element is None:
        element = OxmlElement("w:tblW")
        properties.insert(0, element)
    element.set(qn("w:type"), "dxa")
    element.set(qn("w:w"), str(round(total * 1440)))
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def add_text(paragraph, text, *, bold=False, italic=False, size=10, color=INK):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return run


def setup(document):
    section = document.sections[0]
    section.page_height = Inches(11)
    section.page_width = Inches(8.5)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.66)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.32)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(5)
    for style_name, size in (("Title", 22), ("Heading 1", 13), ("Heading 2", 10.5)):
        style = document.styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY if style_name == "Title" else TEAL
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(header, "HISTORIA CLÍNICA  /  CLINICAL HISTORY    •    CONFIDENCIAL", size=7.5, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(footer, "Formulario del paciente  /  Patient-completed form     ·     ", size=7.5, color=GRAY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def section(document, number, spanish, english, note=None):
    document.add_heading(f"{number}. {spanish} / {english}", level=1)
    if note:
        p = document.add_paragraph(style="Normal")
        add_text(p, note, italic=True, size=8.5, color=GRAY)


def prompt(document, spanish, english, *, rows=1):
    table = document.add_table(rows=1, cols=2)
    set_table_width(table, [2.37, 4.59])
    label, answer = table.rows[0].cells
    label.width = Inches(2.37)
    answer.width = Inches(4.59)
    label.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    question = label.paragraphs[0]
    question.paragraph_format.space_after = Pt(0)
    add_text(question, spanish, bold=True, size=9)
    question.add_run("\n")
    add_text(question, english, italic=True, size=8.2, color=GRAY)
    shade(answer)
    entry = answer.paragraphs[0]
    entry.paragraph_format.space_after = Pt(3)
    add_text(entry, "[Escriba aquí / Type here]", italic=True, size=9, color=GRAY)
    for _ in range(rows - 1):
        answer.add_paragraph(" ")
    for cell in (label, answer):
        set_cell_border(cell)
    no_split(table.rows[0])
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def grid(document, headers, widths, count, *, hint="[Completar / Fill in]"):
    table = document.add_table(rows=count + 1, cols=len(headers))
    set_table_width(table, widths)
    for column, ((spanish, english), width) in enumerate(zip(headers, widths)):
        cell = table.cell(0, column)
        cell.width = Inches(width)
        shade(cell, "DDE9F1")
        p = cell.paragraphs[0]
        add_text(p, spanish, bold=True, size=8.5, color=NAVY)
        p.add_run("\n")
        add_text(p, english, italic=True, size=7.5, color=GRAY)
        set_cell_border(cell)
    repeating_header(table.rows[0])
    no_split(table.rows[0])
    for row in table.rows[1:]:
        no_split(row)
        for column, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = Inches(width)
            shade(cell, FILL if column % 2 == 0 else "FFFFFF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(13)
            add_text(p, hint, italic=True, size=7.5, color=GRAY)
            set_cell_border(cell)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def checklist(document, rows):
    table = document.add_table(rows=1, cols=2)
    set_table_width(table, [4.15, 2.81])
    cells = table.rows[0].cells
    for cell, label in zip(cells, ("Síntoma / Symptom", "Sí / Yes    No    No sé / Unsure")):
        shade(cell, "DDE9F1")
        add_text(cell.paragraphs[0], label, bold=True, size=8.5, color=NAVY)
        set_cell_border(cell)
    repeating_header(table.rows[0])
    for spanish, english in rows:
        cells = table.add_row().cells
        p = cells[0].paragraphs[0]
        add_text(p, spanish, size=9)
        p.add_run("\n")
        add_text(p, english, italic=True, size=8, color=GRAY)
        add_text(cells[1].paragraphs[0], "☐ Sí    ☐ No    ☐ No sé", size=9)
        for cell in cells:
            set_cell_border(cell)
        no_split(table.rows[-1])
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build(output):
    document = Document()
    setup(document)
    document.add_heading("Historia clínica para completar", level=0)
    subtitle = document.add_paragraph()
    add_text(subtitle, "Patient-completed clinical history", size=13, color=TEAL)
    document.add_paragraph("Para conversar con sus profesionales de salud y organizar información que puede faltar en los informes. / For discussion with your clinicians and to organize information that may not appear in reports.")
    p = document.add_paragraph()
    add_text(p, "Cómo llenarlo / How to complete: ", bold=True, size=9)
    add_text(p, "Escriba directamente en las celdas; añada filas si necesita. Si no recuerda un dato, escriba «no sé»; si una pregunta no corresponde, «no aplica». Marque fechas aproximadas como tales. Puede contestar en español o inglés. / Type in the cells; add rows if needed. Use “unsure” or “not applicable” rather than guessing. Mark approximate dates. Answer in either language.", size=9)
    p = document.add_paragraph()
    add_text(p, "Si dispone de poco tiempo / If time is limited: ", bold=True, size=9)
    add_text(p, "Empiece por las secciones 1, 2, 3, 6 y 7; el resto puede completarlo después. / Start with sections 1, 2, 3, 6, and 7; you can fill in the rest later.", size=9)
    p = document.add_paragraph()
    add_text(p, "Privacidad / Privacy: ", bold=True, size=9)
    add_text(p, "Este archivo contiene datos de salud. Envíelo solo a personas de confianza por un medio seguro; puede omitir preguntas que prefiera comentar directamente. No incluya números de identificación innecesarios. / This file contains health information. Share it securely with trusted recipients; you may leave sensitive questions for a private conversation. Avoid unnecessary identification numbers.", size=9)
    p = document.add_paragraph()
    add_text(p, "Si tiene síntomas graves actuales (por ejemplo, dolor intenso en el pecho, dificultad importante para respirar, nueva debilidad o confusión), busque atención urgente; no espere a terminar este formulario. / For severe current symptoms, seek urgent medical attention rather than waiting to finish this form.", bold=True, size=8.5, color=NAVY)

    section(document, 1, "Datos básicos y motivo de consulta", "Basics and main concern")
    prompt(document, "Nombre que desea usar", "Name you prefer to use")
    prompt(document, "Fecha en que completa el formulario; edad actual", "Date completed; current age")
    prompt(document, "¿Quién responde? (paciente / familiar / cuidador y relación)", "Who is answering? (patient / family / caregiver and relationship)")
    prompt(document, "¿Cuál es el principal problema de salud que quiere aclarar?", "What is the main health question you want clarified?", rows=2)
    prompt(document, "¿Qué le preocupa más? ¿Qué espera obtener de la consulta?", "What worries you most? What would you like from this review?", rows=2)
    prompt(document, "¿Qué le han explicado sus médicos hasta ahora? Distinga diagnóstico confirmado de sospecha.", "What have your clinicians said so far? Separate confirmed diagnoses from possibilities.", rows=2)

    section(document, 2, "Síntomas actuales", "Current symptoms", "Si tiene más de tres, copie el bloque o añádalos al final. / Copy this block or add notes for more than three symptoms.")
    for number in range(1, 4):
        document.add_heading(f"Síntoma {number} / Symptom {number}", level=2)
        prompt(document, "¿Qué siente y dónde? (incluya el lado del cuerpo)", "What happens and where? Include which side of the body.")
        prompt(document, "¿Cuándo empezó? Fecha aproximada y si comenzó de repente o poco a poco", "When did it start? Approximate date; sudden or gradual onset.")
        prompt(document, "Frecuencia, duración, intensidad (0–10) y cambios con el tiempo", "Frequency, duration, severity (0–10), and how it has changed.")
        prompt(document, "¿Qué lo empeora o mejora? ¿Se relaciona con alimentos, actividad, horario o medicamentos?", "What makes it worse or better? Relation to food, activity, time of day, or medication.")
        prompt(document, "¿Cómo afecta trabajo, sueño, ejercicio o actividades diarias?", "Effect on work, sleep, exercise, and daily activities.")

    section(document, 3, "Cronología", "Timeline", "No deduzca la fecha a partir del nombre de un archivo; use lo que recuerda o indique «aprox.». / Use recalled or documented dates; mark approximations.")
    grid(document,
         [("Fecha/edad", "Date/age"), ("Síntoma o evento", "Symptom/event"), ("Prueba o tratamiento", "Test/treatment"), ("Qué pasó después", "What happened next")],
         [1.1, 1.95, 1.95, 1.96], 7)

    section(document, 4, "Preguntas breves por sistemas", "Symptom checklist", "Marque una opción en cada fila; si marca «Sí», añada detalles en la sección 2 o al final. / Check one per row; add details for “Yes” answers.")
    checklist(document, [
        ("Fiebre, sudoración nocturna o pérdida de peso sin intención", "Fever, night sweats, or unintended weight loss"),
        ("Cansancio importante o debilidad", "Marked fatigue or weakness"),
        ("Dolor de pecho, palpitaciones, desmayo o hinchazón", "Chest pain, palpitations, fainting, or swelling"),
        ("Falta de aire, tos persistente o silbidos", "Breathlessness, persistent cough, or wheezing"),
        ("Dolor abdominal, vómitos, cambios en evacuaciones o sangre", "Abdominal pain, vomiting, bowel changes, or bleeding"),
        ("Cambios al orinar, sangre en orina o sed excesiva", "Urinary changes, blood in urine, or excessive thirst"),
        ("Dolor o inflamación de articulaciones, músculos o huesos", "Joint, muscle, or bone pain or swelling"),
        ("Sarpullidos, lesiones, cambios en piel, boca u ojos", "Rashes, lesions, or skin, mouth, or eye changes"),
        ("Dolor de cabeza, mareo, entumecimiento o cambios de memoria", "Headache, dizziness, numbness, or memory changes"),
        ("Problemas de sueño, ánimo o concentración", "Sleep, mood, or concentration problems"),
        ("Infecciones repetidas o reacciones alérgicas", "Repeated infections or allergic reactions"),
    ])

    section(document, 5, "Antecedentes médicos", "Medical background")
    grid(document,
         [("Problema/diagnóstico", "Condition/diagnosis"), ("Confirmado o sospechado", "Confirmed or suspected"), ("Año y profesional", "Year and clinician"), ("Tratamiento/estado actual", "Treatment/current status")],
         [2.0, 1.45, 1.55, 1.96], 5)
    prompt(document, "Operaciones, hospitalizaciones o urgencias (cuándo y motivo)", "Surgeries, hospital stays, or emergency visits (when and why)", rows=2)
    prompt(document, "Diagnósticos descartados y por qué (si lo sabe)", "Diagnoses ruled out and why, if known", rows=2)

    section(document, 6, "Medicamentos y alergias", "Medications and allergies", "Incluya los medicamentos actuales, recientes y suplementos. No suspenda medicamentos para completar este formulario. / Include current and recent medicines and supplements; do not stop medicines to complete the form.")
    grid(document,
         [("Nombre", "Name"), ("Dosis/frecuencia", "Dose/frequency"), ("Desde/hasta", "Start/stop dates"), ("Para qué; efecto o reacción", "Reason; benefit/reaction")],
         [1.65, 1.55, 1.4, 2.36], 6)
    prompt(document, "Alergias a medicamentos, alimentos o sustancias: ¿qué reacción tuvo y cuándo?", "Medication, food, or other allergies: what reaction and when?", rows=2)
    prompt(document, "¿Cambió algún medicamento cerca del inicio de los síntomas?", "Did any medication change near symptom onset?")

    section(document, 7, "Pruebas y especialistas", "Tests and specialists")
    grid(document,
         [("Fecha", "Date"), ("Prueba / especialista", "Test / specialist"), ("Resultado según informe", "Result per report"), ("¿Adjunta informe?", "Report attached?")],
         [1.15, 2.0, 2.45, 1.36], 6)
    prompt(document, "¿Hay estudios pendientes o consultas programadas? ¿Cuándo?", "Any pending tests or appointments? When?")
    prompt(document, "¿Tiene análisis, imágenes, biopsias o resultados genéticos que todavía no ha compartido?", "Any laboratory, imaging, biopsy, or genetic reports not yet shared?")

    section(document, 8, "Historia familiar", "Family history", "Indique parentesco, problema y edad aproximada al diagnóstico; «no sé» también es útil. / Include relationship, condition, and approximate age at diagnosis; “unsure” is useful.")
    grid(document,
         [("Parentesco", "Relationship"), ("Enfermedad o síntoma", "Condition/symptom"), ("Edad al inicio", "Age at onset"), ("¿Confirmado?", "Confirmed?")],
         [1.48, 2.64, 1.42, 1.42], 4)
    prompt(document, "¿Hay familiares con problemas similares, enfermedades autoinmunes, cáncer temprano o condiciones hereditarias?", "Similar symptoms, autoimmune disease, early cancer, or inherited conditions in relatives?", rows=2)

    section(document, 9, "Hábitos y exposiciones", "Lifestyle and exposures", "Puede omitir lo que prefiera responder en persona. / You may leave sensitive items to discuss privately.")
    prompt(document, "Trabajo, actividades y contacto con químicos, polvo, metales, humo u otros agentes", "Work, activities, and exposures to chemicals, dust, metals, smoke, or other agents")
    prompt(document, "Viajes, agua/alimentos de riesgo, animales, picaduras o contacto con personas enfermas", "Travel, relevant food/water, animals, insect bites, or sick contacts")
    prompt(document, "Tabaco/vapeo, alcohol, cannabis u otras sustancias: actual o anterior, cantidad aproximada", "Tobacco/vaping, alcohol, cannabis, or other substances: current/past and approximate amount")
    prompt(document, "Dieta o suplementos especiales, cambios importantes de peso, sueño y actividad física", "Diet/supplements, major weight changes, sleep, and physical activity")

    section(document, 10, "Salud reproductiva (si corresponde)", "Reproductive health (if relevant)", "Esta sección es opcional. Comparta solo lo pertinente para interpretar síntomas y elegir pruebas o tratamientos seguros. / Optional: share what is relevant to symptoms and safe test/treatment decisions.")
    prompt(document, "Ciclos menstruales: regulares/irregulares, última menstruación si la sabe, sangrado inusual", "Menstrual cycles: regular/irregular, last period if known, unusual bleeding")
    prompt(document, "Embarazos previos, embarazo posible/actual o planes de embarazo (si aplica)", "Past pregnancies, possible/current pregnancy, or pregnancy plans, if applicable")
    prompt(document, "Anticoncepción, tratamiento hormonal o síntomas de perimenopausia/menopausia (si aplica)", "Contraception, hormone treatment, or perimenopause/menopause symptoms, if applicable")

    section(document, 11, "Impacto, prioridades y notas", "Impact, priorities and notes")
    prompt(document, "¿Qué tareas ha dejado de hacer? ¿Qué días son mejores o peores?", "Which activities are harder? What distinguishes better and worse days?", rows=2)
    prompt(document, "¿Qué tratamientos o cambios ha probado? ¿Qué ayudó, qué no, y qué causó problemas?", "What treatments or changes have you tried? What helped, failed, or caused problems?", rows=2)
    prompt(document, "Sus tres preguntas principales para el médico", "Your top three questions for the clinician", rows=3)
    prompt(document, "Otra información importante que no preguntamos; datos inciertos que quiere verificar", "Other relevant information or uncertain details you want checked", rows=3)
    document.add_paragraph("Gracias por completar lo que pueda. Revisaremos cualquier dato citado con el informe original; sus respuestas son información reportada por usted hasta que su equipo clínico las confirme. / Thank you for sharing what you can. Document-derived values must be checked against the original report; your answers remain patient-reported until confirmed by your clinical team.")
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return document


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path,
                        default=Path(__file__).resolve().parents[1] / "clinical_history_questionnaire.docx")
    args = parser.parse_args()
    build(args.output)
    print("Created", args.output)


if __name__ == "__main__":
    main()
