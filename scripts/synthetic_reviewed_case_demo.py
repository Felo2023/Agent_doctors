"""Create a source-linked, human-reviewed fictional case for bus integration.

All names, dates and findings are fabricated. The output is explicitly marked
reviewed_subset and NOT ready for diagnosis.
"""

from pathlib import Path
import tempfile

from docx import Document
from openpyxl import Workbook

from case_review.store import ReviewStore
from intake.extract import extract_snapshot
from intake.markdown import write_bundle
from scripts.synthetic_intake_demo import create_page


def build(case_id="SYNTH-002"):
    root = Path(tempfile.mkdtemp(prefix="synthetic-reviewed-case-"))
    root.chmod(0o700)
    (root / ".synthetic-fixture").write_text("SYNTHETIC_INTAKE_FIXTURE_V1", encoding="utf-8")
    pdf = root / "fictional-result.pdf"
    create_page(pdf, "1,4", scan=False)
    word = root / "fictional-history.docx"
    document = Document()
    document.add_paragraph("In this fictional example the patient reports fatigue starting in early 2024.")
    document.save(word)
    excel = root / "fictional-workbook.xlsx"
    workbook = Workbook()
    workbook.active.title = "Tests"
    workbook.active["A1"] = "Fictional follow-up requested"
    workbook.save(excel)

    snapshot = extract_snapshot([pdf, word, excel], root / "extraction.json",
                                fixture_root=root, case_id=case_id)
    write_bundle(snapshot, root, fixture_root=root)
    store = ReviewStore(root)
    try:
        candidate = snapshot["documents"][0]["candidates"][0]
        store.decide(candidate["id"], "verify", "fictional-reviewer-B",
                     "Compared fictional lab line against original PDF", 0)
        history_doc = next(doc for doc in snapshot["documents"] if doc["format"] == "docx")
        fact = store.propose_fact(history_doc["source_id"], history_doc["blocks"][0]["id"], "symptom",
                                  "Fatigue since early 2024 (patient-reported in fictional document)",
                                  "2024", "approximate", "patient_reported", "fictional-proposer-A", 1)
        store.decide_fact(fact["fact_id"], "verify", "fictional-reviewer-B",
                          "Compared fictional statement with Word paragraph", 2)
        excel_doc = next(doc for doc in snapshot["documents"] if doc["format"] == "xlsx")
        item = store.propose_fact(excel_doc["source_id"], excel_doc["blocks"][0]["id"], "test_report",
                                  "Fictional follow-up requested (per workbook)",
                                  None, "unknown", "source_text", "fictional-proposer-A", 3)
        store.decide_fact(item["fact_id"], "verify", "fictional-reviewer-B",
                          "Verified workbook cell A1 against original", 4)
        version = store.commit(5)
        return root, version
    finally:
        store.close()


def main():
    root, version = build()
    print("Synthetic fixture:", root)
    print("Case:", version["case_id"], "version:", version["version"])
    print("Reviewed lab observations:", len(version["observations"]))
    print("Reviewed source statements:", len(version["source_statements"]))
    print("Timeline entries:", len(version["timeline"]))
    print("Extracted-block coverage complete:", version["intake_coverage_complete"])
    print("Ready for diagnosis:", version["ready_for_diagnosis"])


if __name__ == "__main__":
    main()
